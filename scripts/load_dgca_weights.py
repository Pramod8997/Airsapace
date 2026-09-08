"""Compute route weights from DGCA DOM city-pair passenger data (July 2026).

    .venv/bin/python scripts/load_dgca_weights.py [--refetch]

Parses data/fixtures/dgca_citypair_jul2026.xlsx and writes
data/fixtures/dgca_citypair_weights.json. Weights are each route's share of
total basket-route passengers (directional), so they sum to 1.0.
--refetch re-downloads the XLSX from DGCA's public S3 first.
"""
from __future__ import annotations

import json
import sys
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scripts.generate_replay_data import ROUTES  # basket route ids

DGCA_URL = (
    "https://public-prd-dgca.s3.ap-south-1.amazonaws.com/InventoryList/"
    "dataReports/aviationDataStatistics/airTransport/domestic/airTraffic/"
    "monthly/DOM CITYPAIR DATA, JULY 2026.xlsx"
)
XLSX_PATH = Path("data/fixtures/dgca_citypair_jul2026.xlsx")
OUT_PATH = Path("data/fixtures/dgca_citypair_weights.json")
SOURCE_LABEL = "DGCA DOM city-pair data July 2026"

# DGCA city names in the sheet -> basket IATA codes. Mumbai is split across
# two airport rows (Mumbai + Navi Mumbai); both roll up into the city pair.
CITY_TO_IATA = {
    "delhi": "DEL",
    "bengaluru": "BLR",
    "bangalore": "BLR",
    "mumbai (mumbai)": "BOM",
    "mumbai (navi mumbai)": "BOM",
    "kolkata": "CCU",
    "chennai": "MAA",
    "hyderabad": "HYD",
}

# Sheet layout: header on row 3 (1-based); columns 2/3 are CITY 1 / CITY 2,
# column 4 "PASSENGERS TO CITY 2" (city1 -> city2), column 5
# "PASSENGERS FROM CITY 2" (city2 -> city1).
HEADER_ROW = 3
COL_CITY1, COL_CITY2, COL_TO, COL_FROM = 2, 3, 4, 5


def _norm(name: object) -> str:
    return " ".join(str(name).strip().lower().split())


def route_pax_from_workbook(path: Path) -> dict[str, int]:
    """Directional passenger counts for all basket-city pairs, keyed "DEL-BOM".

    Rows are unordered city pairs with two directional passenger columns, so
    each row contributes to two route keys. Multiple rows for one city pair
    (Mumbai's two airports) are summed.
    """
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    pax: dict[str, int] = {}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        c1, c2 = _norm(row[COL_CITY1 - 1]), _norm(row[COL_CITY2 - 1])
        a, b = CITY_TO_IATA.get(c1), CITY_TO_IATA.get(c2)
        to_c2, from_c2 = row[COL_TO - 1], row[COL_FROM - 1]
        if a is None or b is None or not to_c2 or not from_c2:
            continue
        pax[f"{a}-{b}"] = pax.get(f"{a}-{b}", 0) + int(to_c2)
        pax[f"{b}-{a}"] = pax.get(f"{b}-{a}", 0) + int(from_c2)
    return pax


def compute_weights(route_pax: dict[str, int]) -> dict[str, float]:
    """Route share of total basket-route passengers, normalized over the 10
    basket routes only."""
    basket_pax = {r: route_pax[r] for r in ROUTES}
    missing = [r for r, p in basket_pax.items() if p <= 0]
    if missing:
        raise ValueError(f"no DGCA passenger data for basket routes: {missing}")
    total = sum(basket_pax.values())
    return {r: round(p / total, 6) for r, p in basket_pax.items()}


def write_fixture(out_path: Path = OUT_PATH) -> dict:
    route_pax = route_pax_from_workbook(XLSX_PATH)
    weights = compute_weights(route_pax)
    payload = {
        "source": SOURCE_LABEL,
        "url": DGCA_URL,
        "route_weights": weights,
        "route_pax": {r: route_pax[r] for r in ROUTES},
    }
    out_path.write_text(json.dumps(payload, indent=2) + "\n")
    return payload


def refetch() -> None:
    print(f"downloading {DGCA_URL}")
    req = urllib.request.Request(DGCA_URL, headers={"User-Agent": "AirStat-India/0.1"})
    with urllib.request.urlopen(req, timeout=60) as resp, XLSX_PATH.open("wb") as f:
        f.write(resp.read())
    print(f"saved {XLSX_PATH} ({XLSX_PATH.stat().st_size} bytes)")


def main() -> None:
    if "--refetch" in sys.argv:
        refetch()
    payload = write_fixture()
    print(f"wrote {OUT_PATH}\n")
    print(f"{'route':10} {'pax':>9} {'weight':>9}")
    for route, pax in payload["route_pax"].items():
        print(f"{route:10} {pax:>9,} {payload['route_weights'][route]:>9.6f}")
    print(f"{'total':10} {sum(payload['route_pax'].values()):>9,} "
          f"{sum(payload['route_weights'].values()):>9.6f}")


if __name__ == "__main__":
    main()
